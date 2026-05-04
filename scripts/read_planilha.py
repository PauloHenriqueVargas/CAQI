"""
Lê Planilha_Base_CAQi_Preenchida.xlsx (e a base vazia para comparar)
e despeja: nomes de abas, dimensões, headers, conteúdo completo, e fórmulas.
"""
import sys
from openpyxl import load_workbook
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PATHS = [
    ROOT / "docs" / "references" / "Planilha_Base_CAQi_Preenchida.xlsx",
    ROOT / "docs" / "references" / "Planilha_Base_CAQi.xlsx",
]
OUT = ROOT / "scripts" / "planilha_dump.txt"


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def dump(path: Path, w, max_rows: int = 500):
    w.write(f"\n{'='*80}\n")
    w.write(f"ARQUIVO: {path.name}\n")
    w.write(f"{'='*80}\n")

    wb_v = load_workbook(path, data_only=True)
    wb_f = load_workbook(path, data_only=False)

    w.write(f"\nAbas ({len(wb_v.sheetnames)}): {wb_v.sheetnames}\n")

    for sheet_name in wb_v.sheetnames:
        ws_v = wb_v[sheet_name]
        ws_f = wb_f[sheet_name]
        w.write(f"\n{'-'*80}\n")
        w.write(f"ABA: {sheet_name!r}  ({ws_v.max_row} linhas x {ws_v.max_column} colunas)\n")
        w.write(f"{'-'*80}\n")

        for row_idx, (row_v, row_f) in enumerate(zip(ws_v.iter_rows(), ws_f.iter_rows()), start=1):
            if row_idx > max_rows:
                w.write(f"  ... (truncado em {max_rows} linhas)\n")
                break
            cells = []
            for cv, cf in zip(row_v, row_f):
                val = fmt(cv.value)
                if cf.value is not None and isinstance(cf.value, str) and cf.value.startswith("="):
                    val = f"{val} [FORMULA: {cf.value}]"
                cells.append(val)
            if any(c.strip() for c in cells):
                w.write(f"  R{row_idx:3d} | " + " | ".join(cells) + "\n")


if __name__ == "__main__":
    with OUT.open("w", encoding="utf-8") as w:
        for p in PATHS:
            if p.exists():
                dump(p, w)
            else:
                w.write(f"NOT FOUND: {p}\n")
    print(f"OK -> {OUT}")
